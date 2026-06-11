#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Service de surveillance automatique des bilans hebdomadaires.

Surveille un dossier (par défaut : 'bilans/') pour les nouveaux fichiers Excel.
Quand un nouveau bilan arrive (ou qu'un existant est modifié) :
  1. Détecte automatiquement la semaine et l'année depuis le nom de fichier
  2. Lit les données du bilan
  3. Met à jour (ou insère) les données dans le fichier BDD Excel principal
  4. Relance update_dashboard_v2_FINAL.py pour rafraîchir le dashboard HTML
  5. Évite les doublons (même semaine/secteur déjà en BDD = mise à jour)

Usage :
  python watcher_service.py                    # surveille ./bilans/
  python watcher_service.py --folder C:/bilans # surveille un autre dossier
  python watcher_service.py --interval 15      # vérification toutes les 15s (défaut 30s)
  python watcher_service.py --once             # traite les fichiers existants une fois et quitte

Format attendu des fichiers bilan :
  Nom de fichier : Bilan_S23_2026.xlsx, bilan_semaine23_2026.xlsx, S23-2026.xlsx...
  Feuille production : colonnes Semaine, Secteur, Prod_Globale, Vit_prod, etc.
  Feuille RH (optionnel) : colonnes Mois, Annee, Secteur, NB_ETP, Montant_MS, etc.
"""

import os, sys, re, json, time, hashlib, shutil, argparse, logging, subprocess
from datetime import datetime
import pandas as pd
import openpyxl

# ─────────────────────────────────────────────
# CONFIGURATION
# ─────────────────────────────────────────────

DEFAULT_WATCH_FOLDER = 'bilans'
DEFAULT_INTERVAL = 30
STATE_FILE = '.watcher_state.json'
UPDATE_SCRIPT = 'update_dashboard_v2_FINAL.py'
LOG_FILE = 'watcher_service.log'

SEMAINE_MOIS = {}
_s = 1
for _mois, _nb in [('Janvier',5),('Février',4),('Mars',4),('Avril',5),
    ('Mai',4),('Juin',4),('Juillet',5),('Août',4),
    ('Septembre',4),('Octobre',5),('Novembre',4),('Décembre',4)]:
    for _i in range(_nb):
        SEMAINE_MOIS[f'S{_s:02d}'] = _mois
        _s += 1

# ─────────────────────────────────────────────
# LOGGING
# ─────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s  %(levelname)-8s %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S',
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler(LOG_FILE, encoding='utf-8')
    ]
)
log = logging.getLogger('watcher')

# ─────────────────────────────────────────────
# ÉTAT (quels fichiers ont déjà été traités)
# ─────────────────────────────────────────────

def load_state():
    if os.path.exists(STATE_FILE):
        try:
            with open(STATE_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def save_state(state):
    with open(STATE_FILE, 'w', encoding='utf-8') as f:
        json.dump(state, f, indent=2, ensure_ascii=False)

def file_hash(path):
    h = hashlib.md5()
    try:
        with open(path, 'rb') as f:
            while True:
                chunk = f.read(65536)
                if not chunk: break
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return ''

# ─────────────────────────────────────────────
# DÉTECTION SEMAINE / ANNÉE DANS LE NOM DE FICHIER
# ─────────────────────────────────────────────

def detect_semaine_annee(filename):
    """Extrait (semaine_str, annee_int) depuis le nom de fichier."""
    name = os.path.basename(filename)
    patterns = [
        r'[Ss](?:emaine[\s_-]?)?(\d{1,2})[\s_-]+(\d{4})',
        r'[Ss](?:em[\s_-]?)?(\d{1,2})[-_](\d{4})',
        r'(\d{4})[-_][Ss](?:em[\s_-]?)?(\d{1,2})',
        r'[Ss][Ee][Mm][\s_-]?(\d{1,2})[\s_-]+(\d{4})',
        r'week[\s_-]?(\d{1,2})[\s_-]+(\d{4})',
    ]
    for pat in patterns:
        m = re.search(pat, name, re.IGNORECASE)
        if m:
            g = m.groups()
            if len(g) == 2:
                a, b = g
                # Déterminer lequel est la semaine (1-53) et lequel est l'année
                if 1 <= int(a) <= 53 and int(b) > 100:
                    return f'S{int(a):02d}', int(b)
                elif 1 <= int(b) <= 53 and int(a) > 100:
                    return f'S{int(b):02d}', int(a)
    return None, None

# ─────────────────────────────────────────────
# TROUVER LE FICHIER BDD PRINCIPAL
# ─────────────────────────────────────────────

def find_bdd_file():
    """Trouve le fichier Excel BDD dans le dossier courant."""
    candidates = [f for f in os.listdir('.')
                  if f.lower().endswith(('.xlsx', '.xls'))
                  and not f.startswith('~')
                  and any(k in f.lower() for k in ['bdd', 'base', 'database', 'données', 'donnees'])]
    if candidates:
        return sorted(candidates)[-1]
    # Fallback : prendre le seul Excel du dossier
    excels = [f for f in os.listdir('.') if f.lower().endswith(('.xlsx','.xls')) and not f.startswith('~')]
    return sorted(excels)[-1] if excels else None

def find_prod_sheet(wb_or_xl):
    """Trouve le nom de la feuille production."""
    names = wb_or_xl.sheet_names if hasattr(wb_or_xl, 'sheet_names') else wb_or_xl.sheetnames
    for sheet in names:
        n = sheet.lower().replace(' ', '').replace('_', '')
        if any(k in n for k in ['production', 'prod', 'bddprod']):
            return sheet
    return names[0] if names else None

def find_rh_sheet(wb_or_xl):
    names = wb_or_xl.sheet_names if hasattr(wb_or_xl, 'sheet_names') else wb_or_xl.sheetnames
    for sheet in names:
        n = sheet.lower()
        if any(k in n for k in ['rh', 'ressource', 'personnel', 'hr', 'admin', 'effectif']):
            return sheet
    return None

# ─────────────────────────────────────────────
# LECTURE DU BILAN HEBDOMADAIRE
# ─────────────────────────────────────────────

COL_MAP_PROD = {
    'semaine':     ['semaine', 'week', 'sem'],
    'secteur':     ['secteur', 'sector', 'service', 'unite'],
    'prod_globale':['prod_globale', 'production', 'prod', 'quantite', 'qty', 'volume'],
    'cout_total':  ['cout_total', 'cout', 'cost', 'cout_mod', 'mod'],
    'mod_heures_travaillees': ['heures', 'heures_travaillees', 'heures_mod', 'hrs', 'h_trav'],
    'vit_prod':    ['vit_prod', 'vitesse', 'vitesse_prod', 'speed'],
    'cout_pour_1000': ['cout_pour_1000', 'cout1000', 'c1000', 'cout_mille'],
    'hld':         ['hld', 'respect_hld', 'hld_resp'],
    'taux_cobot':  ['taux_cobot', 'cobot', 'taux_cob', 'cobot_rate'],
    'prod_conducteur_matin': ['prod_conducteur', 'conducteur', 'prod_cond'],
}

COL_MAP_RH = {
    'secteur':   ['secteur', 'sector', 'service'],
    'effectif':  ['effectif', 'nb_effectif', 'headcount', 'effectifs'],
    'montant_ms':['montant_ms', 'masse_salariale', 'ms', 'salaires', 'masse_sal'],
    'cout_interim': ['cout_interim', 'interim', 'cout_int', 'interimaires'],
    'nb_etp':    ['nb_etp', 'etp', 'fte'],
    'nb_interimaires': ['nb_interimaires', 'interimaires', 'nb_interim', 'interim_nb'],
    'heures_absence': ['heures_absence', 'absences', 'abs', 'h_abs'],
    'heures_hs': ['heures_hs', 'hs', 'heures_sup', 'h_sup'],
}

def normalize_col(name):
    import unicodedata
    name = str(name).strip().lower()
    name = unicodedata.normalize('NFD', name)
    name = ''.join(c for c in name if unicodedata.category(c) != 'Mn')
    return re.sub(r'[\s_\-/]+', '_', name)

def map_columns(df_cols, mapping):
    """Retourne dict {cible: nom_colonne_dans_df}."""
    norm_cols = {normalize_col(c): c for c in df_cols}
    result = {}
    for target, aliases in mapping.items():
        for alias in aliases:
            norm_alias = normalize_col(alias)
            if norm_alias in norm_cols:
                result[target] = norm_cols[norm_alias]
                break
    return result

def read_bilan_prod(filepath, semaine, annee):
    """Lit la feuille production d'un bilan et retourne une liste de dicts."""
    try:
        xl = pd.ExcelFile(filepath)
        sheet = find_prod_sheet(xl)
        if not sheet:
            log.warning(f"  Aucune feuille production trouvée dans {filepath}")
            return []
        df = pd.read_excel(filepath, sheet_name=sheet)
        df.columns = [str(c) for c in df.columns]
        col_map = map_columns(df.columns, COL_MAP_PROD)

        mois = SEMAINE_MOIS.get(semaine, 'Janvier')
        rows = []
        for _, r in df.iterrows():
            secteur = str(r.get(col_map.get('secteur', ''), 'Inconnu')).strip()
            if not secteur or secteur in ('nan', 'None', ''):
                continue
            row = {
                'Semaine': semaine,
                'Mois': mois,
                'Annee': annee,
                'Secteur': secteur,
                'Prod_Globale': _safe_float(r, col_map.get('prod_globale')),
                'Cout_total': _safe_float(r, col_map.get('cout_total')),
                'MOD_heures_travaillees': _safe_float(r, col_map.get('mod_heures_travaillees')),
                'Vit_prod': _safe_float(r, col_map.get('vit_prod')),
                'Cout_pour_1000': _safe_float(r, col_map.get('cout_pour_1000')),
                'HLD': _safe_float(r, col_map.get('hld')),
                'Taux_COBOT': _safe_float(r, col_map.get('taux_cobot')),
                'Prod_Conducteur_Matin': _safe_float(r, col_map.get('prod_conducteur_matin')),
            }
            rows.append(row)
        log.info(f"  Bilan prod : {len(rows)} lignes lues depuis '{sheet}'")
        return rows
    except Exception as e:
        log.error(f"  Erreur lecture bilan prod {filepath}: {e}")
        return []

def read_bilan_rh(filepath, semaine, annee):
    """Lit la feuille RH d'un bilan (optionnel)."""
    try:
        xl = pd.ExcelFile(filepath)
        sheet = find_rh_sheet(xl)
        if not sheet:
            return []
        df = pd.read_excel(filepath, sheet_name=sheet)
        df.columns = [str(c) for c in df.columns]
        col_map = map_columns(df.columns, COL_MAP_RH)
        mois = SEMAINE_MOIS.get(semaine, 'Janvier')
        rows = []
        for _, r in df.iterrows():
            secteur = str(r.get(col_map.get('secteur', ''), 'Inconnu')).strip()
            if not secteur or secteur in ('nan', 'None', ''):
                continue
            rows.append({
                'Mois': mois, 'Annee': annee, 'Secteur': secteur,
                'Effectif': _safe_int(r, col_map.get('effectif')),
                'Montant_MS': _safe_float(r, col_map.get('montant_ms')),
                'Cout_interim': _safe_float(r, col_map.get('cout_interim')),
                'NB_ETP': _safe_float(r, col_map.get('nb_etp')),
                'NB_Interimaires': _safe_float(r, col_map.get('nb_interimaires')),
                'Heures_Absence': _safe_float(r, col_map.get('heures_absence')),
                'Heures_HS': _safe_float(r, col_map.get('heures_hs')),
            })
        log.info(f"  Bilan RH : {len(rows)} lignes lues depuis '{sheet}'")
        return rows
    except Exception as e:
        log.warning(f"  Lecture RH ignorée ({e})")
        return []

def _safe_float(row, col):
    if not col or col not in row.index: return None
    v = row[col]
    try: return float(v) if pd.notna(v) else None
    except: return None

def _safe_int(row, col):
    v = _safe_float(row, col)
    return int(v) if v is not None else None

# ─────────────────────────────────────────────
# MISE À JOUR DE LA BDD EXCEL
# ─────────────────────────────────────────────

def update_bdd_prod(bdd_file, new_rows, semaine, annee):
    """Ajoute ou remplace les lignes de la semaine dans la BDD production."""
    if not new_rows:
        return 0

    xl = pd.ExcelFile(bdd_file)
    sheet = find_prod_sheet(xl)
    if not sheet:
        log.error(f"  Feuille production introuvable dans {bdd_file}")
        return 0

    df = pd.read_excel(bdd_file, sheet_name=sheet)

    # Supprimer les lignes existantes pour cette semaine/année (évite les doublons)
    mask = (df['Semaine'].astype(str) == str(semaine)) & (df['Annee'].astype(str) == str(annee))
    n_existing = mask.sum()
    df = df[~mask]

    # Ajouter les nouvelles lignes
    new_df = pd.DataFrame(new_rows)
    # Aligner les colonnes
    for col in df.columns:
        if col not in new_df.columns:
            new_df[col] = None
    new_df = new_df.reindex(columns=df.columns)
    df = pd.concat([df, new_df], ignore_index=True)
    df = df.sort_values(['Annee', 'Semaine', 'Secteur'], na_position='last')

    # Sauvegarder
    _save_sheet(bdd_file, sheet, df)
    action = 'mis à jour' if n_existing > 0 else 'ajoutés'
    log.info(f"  BDD prod : {len(new_rows)} lignes {action} pour {semaine}/{annee} (remplacé {n_existing})")
    return len(new_rows)

def update_bdd_rh(bdd_file, new_rows, semaine, annee):
    """Ajoute ou remplace les lignes RH du mois correspondant."""
    if not new_rows:
        return 0
    mois = SEMAINE_MOIS.get(semaine, 'Janvier')

    xl = pd.ExcelFile(bdd_file)
    sheet = find_rh_sheet(xl)
    if not sheet:
        log.warning(f"  Feuille RH introuvable dans {bdd_file} — RH ignoré")
        return 0

    df = pd.read_excel(bdd_file, sheet_name=sheet)
    # RH : on met à jour par mois/année/secteur
    for row in new_rows:
        mask = ((df.get('Mois','') == mois) &
                (df.get('Annee', 0).astype(str) == str(annee)) &
                (df.get('Secteur', '') == row['Secteur']))
        df = df[~mask]

    new_df = pd.DataFrame(new_rows)
    for col in df.columns:
        if col not in new_df.columns:
            new_df[col] = None
    new_df = new_df.reindex(columns=df.columns)
    df = pd.concat([df, new_df], ignore_index=True)
    _save_sheet(bdd_file, sheet, df)
    log.info(f"  BDD RH : {len(new_rows)} lignes mises à jour pour {mois}/{annee}")
    return len(new_rows)

def _save_sheet(bdd_file, sheet_name, df):
    """Écrit un DataFrame dans une feuille Excel en préservant les autres feuilles."""
    wb = openpyxl.load_workbook(bdd_file)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    # En-têtes
    ws.append(list(df.columns))
    # Données
    for _, row in df.iterrows():
        ws.append([None if pd.isna(v) else v for v in row])
    wb.save(bdd_file)
    wb.close()

# ─────────────────────────────────────────────
# TRAITEMENT D'UN FICHIER BILAN
# ─────────────────────────────────────────────

def process_bilan(filepath, bdd_file):
    """Traite un fichier bilan : lit, met à jour BDD, relance le script."""
    filename = os.path.basename(filepath)
    semaine, annee = detect_semaine_annee(filename)

    if not semaine or not annee:
        log.warning(f"  {filename} — Impossible de détecter la semaine/année. "
                    f"Nommez le fichier comme : Bilan_S23_2026.xlsx")
        return False

    log.info(f"  Traitement : {filename} → {semaine} / {annee}")

    # Sauvegarde préventive de la BDD
    backup = bdd_file + f'.bak_{semaine}_{annee}'
    try:
        shutil.copy2(bdd_file, backup)
    except Exception:
        pass

    try:
        prod_rows = read_bilan_prod(filepath, semaine, annee)
        rh_rows = read_bilan_rh(filepath, semaine, annee)

        n_prod = update_bdd_prod(bdd_file, prod_rows, semaine, annee)
        n_rh = update_bdd_rh(bdd_file, rh_rows, semaine, annee)

        if n_prod + n_rh == 0:
            log.warning(f"  {filename} — Aucune donnée extraite (vérifiez le format)")
            return False

        # Nettoyer la sauvegarde si succès
        try: os.remove(backup)
        except Exception: pass

        # Relancer le script de mise à jour du dashboard
        run_update_script()
        return True

    except Exception as e:
        log.error(f"  Erreur traitement {filename}: {e}")
        # Restaurer la sauvegarde en cas d'erreur
        try:
            shutil.copy2(backup, bdd_file)
            log.info(f"  BDD restaurée depuis sauvegarde")
        except Exception:
            pass
        return False

def run_update_script():
    """Exécute update_dashboard_v2_FINAL.py."""
    if not os.path.exists(UPDATE_SCRIPT):
        log.warning(f"  {UPDATE_SCRIPT} introuvable — dashboard non mis à jour")
        return
    try:
        result = subprocess.run(
            [sys.executable, UPDATE_SCRIPT],
            capture_output=True, text=True, timeout=120
        )
        if result.returncode == 0:
            log.info("  ✅ Dashboard mis à jour avec succès")
        else:
            log.error(f"  ❌ Erreur mise à jour dashboard:\n{result.stderr[:500]}")
    except subprocess.TimeoutExpired:
        log.error("  ❌ Timeout lors de la mise à jour du dashboard")
    except Exception as e:
        log.error(f"  ❌ Impossible de lancer {UPDATE_SCRIPT}: {e}")

# ─────────────────────────────────────────────
# BOUCLE PRINCIPALE DE SURVEILLANCE
# ─────────────────────────────────────────────

def scan_and_process(watch_folder, bdd_file, state, run_once=False):
    """Scanne le dossier et traite les fichiers nouveaux ou modifiés."""
    if not os.path.exists(watch_folder):
        os.makedirs(watch_folder, exist_ok=True)
        log.info(f"Dossier de surveillance créé : {watch_folder}")
        log.info(f"Déposez vos bilans Excel dans ce dossier (ex: Bilan_S23_2026.xlsx)")
        return False

    excel_files = [
        os.path.join(watch_folder, f)
        for f in os.listdir(watch_folder)
        if f.lower().endswith(('.xlsx', '.xls')) and not f.startswith('~')
    ]

    processed_any = False
    for filepath in sorted(excel_files):
        h = file_hash(filepath)
        key = os.path.basename(filepath)
        if state.get(key) == h:
            continue  # Pas de changement

        log.info(f"Nouveau fichier ou modification détectée : {key}")
        success = process_bilan(filepath, bdd_file)
        if success:
            state[key] = h
            processed_any = True
        elif run_once:
            state[key] = h  # Marquer quand même pour ne pas reboucler

    return processed_any

def main():
    parser = argparse.ArgumentParser(description='Service de surveillance des bilans KPI')
    parser.add_argument('--folder', default=DEFAULT_WATCH_FOLDER,
                        help=f'Dossier à surveiller (défaut: {DEFAULT_WATCH_FOLDER})')
    parser.add_argument('--interval', type=int, default=DEFAULT_INTERVAL,
                        help=f'Intervalle de vérification en secondes (défaut: {DEFAULT_INTERVAL})')
    parser.add_argument('--once', action='store_true',
                        help='Traite les fichiers existants une fois et quitte')
    args = parser.parse_args()

    log.info('=' * 60)
    log.info('  Service de surveillance KPI démarré')
    log.info(f'  Dossier surveillé : {os.path.abspath(args.folder)}')
    log.info(f'  Intervalle : {args.interval}s')
    log.info('=' * 60)

    bdd_file = find_bdd_file()
    if not bdd_file:
        log.error("Aucun fichier Excel BDD trouvé dans le dossier courant.")
        log.error("Lancez ce script depuis le même dossier que votre BDD Excel.")
        sys.exit(1)
    log.info(f"Fichier BDD détecté : {bdd_file}")

    state = load_state()

    if args.once:
        scan_and_process(args.folder, bdd_file, state, run_once=True)
        save_state(state)
        log.info("Traitement unique terminé.")
        return

    log.info("Surveillance active — appuyez sur Ctrl+C pour arrêter")
    try:
        while True:
            try:
                changed = scan_and_process(args.folder, bdd_file, state)
                if changed:
                    save_state(state)
            except Exception as e:
                log.error(f"Erreur dans la boucle de surveillance: {e}")
            time.sleep(args.interval)
    except KeyboardInterrupt:
        log.info("Service arrêté par l'utilisateur.")
        save_state(state)

if __name__ == '__main__':
    main()
