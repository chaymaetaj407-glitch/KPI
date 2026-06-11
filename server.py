#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import http.server, socketserver, subprocess
import threading, os, sys, json, webbrowser
from urllib.parse import urlparse
import openpyxl

PORT = 8765
SCRIPT = "update_dashboard_v2_FINAL.py"

# ── Dépôt GitHub pour l'import automatique ──
GITHUB_REPO_OWNER = "chaymaetaj407-glitch"
GITHUB_REPO_NAME  = "KPI"
GITHUB_BRANCH     = "main"

class DashboardHandler(http.server.SimpleHTTPRequestHandler):

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def do_POST(self):
        parsed = urlparse(self.path)

        # ── Recevoir et sauvegarder les commentaires avant mise à jour ──
        if parsed.path == '/api/save-comments':
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length)
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            try:
                comments = json.loads(body.decode('utf-8'))
                # Trouver le fichier Excel
                excel_files = [f for f in os.listdir('.')
                               if f.endswith(('.xlsx','.xls')) and not f.startswith('~')]
                if not excel_files:
                    self.wfile.write(json.dumps({'success': False,
                        'message': 'Aucun fichier Excel trouvé'}).encode())
                    return
                excel_file = sorted(excel_files)[-1]
                # Sauvegarder dans la feuille Commentaires
                wb = openpyxl.load_workbook(excel_file)
                if 'Commentaires' in wb.sheetnames:
                    del wb['Commentaires']
                ws = wb.create_sheet('Commentaires')
                ws.append(['Cle', 'Type', 'Texte', 'Date_modif'])
                for key, data in comments.items():
                    ws.append([key, data.get('type',''),
                               data.get('texte',''), data.get('date','')])
                wb.save(excel_file)
                wb.close()
                self.wfile.write(json.dumps({
                    'success': True,
                    'message': f"{len(comments)} commentaires sauvegardés dans {excel_file}"
                }).encode('utf-8'))
            except Exception as e:
                self.wfile.write(json.dumps({'success': False, 'message': str(e)}).encode())
            return

    def do_GET(self):
        parsed = urlparse(self.path)

        # ── Lancer la mise à jour (+ import GitHub auto) ──
        if parsed.path == '/api/update':
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()

            # Étape 0 : télécharger les nouveaux bilans depuis GitHub
            github_log = sync_github_bilans()

            if not os.path.exists(SCRIPT):
                self.wfile.write(json.dumps({
                    'success': False,
                    'message': f"Script {SCRIPT} introuvable. {github_log}"
                }).encode('utf-8'))
                return
            try:
                result = subprocess.run(
                    [sys.executable, SCRIPT],
                    capture_output=True, text=True,
                    input="\n", timeout=120, encoding='utf-8'
                )
                success = "REUSSIE" in result.stdout or result.returncode == 0
                output = github_log + "\n" + result.stdout + result.stderr
                self.wfile.write(json.dumps({
                    'success': success,
                    'message': output[-3000:] if len(output) > 3000 else output
                }).encode('utf-8'))
            except subprocess.TimeoutExpired:
                self.wfile.write(json.dumps({
                    'success': False, 'message': "Délai dépassé (>2 min)."
                }).encode('utf-8'))
            except Exception as e:
                self.wfile.write(json.dumps({'success': False, 'message': str(e)}).encode('utf-8'))
            return

        # ── Liste des fichiers GitHub disponibles ──
        if parsed.path == '/api/github-files':
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            try:
                files = list_github_bilans()
                self.wfile.write(json.dumps({'success': True, 'files': files}).encode('utf-8'))
            except Exception as e:
                self.wfile.write(json.dumps({'success': False, 'message': str(e)}).encode('utf-8'))
            return

        # ── Statut ──
        if parsed.path == '/api/status':
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(json.dumps({'running': True}).encode('utf-8'))
            return

        super().do_GET()

    def log_message(self, format, *args):
        if args[1] not in ('200', '304'):
            print(f"[Serveur] {format % args}")


# ─────────────────────────────────────────────────────
# IMPORT AUTOMATIQUE GITHUB
# ─────────────────────────────────────────────────────

def list_github_bilans():
    """Liste les fichiers bilan_prod_semaine sur GitHub."""
    import urllib.request
    url = (f"https://api.github.com/repos/{GITHUB_REPO_OWNER}/{GITHUB_REPO_NAME}"
           f"/contents/?ref={GITHUB_BRANCH}")
    req = urllib.request.Request(url, headers={'User-Agent': 'KPI-Dashboard'})
    with urllib.request.urlopen(req, timeout=15) as r:
        items = json.loads(r.read().decode())
    bilan_files = []
    keywords = ['bilan', 'prod', 'semaine']
    bdd_keywords = ['bdd', 'kpi', 'dashboard']
    for item in items:
        name = item.get('name', '')
        if not name.endswith(('.xlsx', '.xls')):
            continue
        fname_low = name.lower()
        if any(k in fname_low for k in bdd_keywords):
            continue
        if any(k in fname_low for k in keywords):
            bilan_files.append({
                'name': name,
                'download_url': item.get('download_url'),
                'sha': item.get('sha')
            })
    return bilan_files


def sync_github_bilans():
    """
    Télécharge depuis GitHub les fichiers bilan non encore présents localement.
    Retourne un log en texte.
    """
    log = []
    try:
        remote_files = list_github_bilans()
        if not remote_files:
            return "[GitHub] Aucun nouveau fichier bilan trouvé."
        log.append(f"[GitHub] {len(remote_files)} fichier(s) bilan détecté(s) sur GitHub.")
        import urllib.request
        for rf in remote_files:
            local_path = rf['name']
            if os.path.exists(local_path):
                log.append(f"  → {rf['name']} : déjà présent localement, ignoré.")
                continue
            url = rf.get('download_url')
            if not url:
                log.append(f"  → {rf['name']} : URL de téléchargement manquante.")
                continue
            try:
                req = urllib.request.Request(url, headers={'User-Agent': 'KPI-Dashboard'})
                with urllib.request.urlopen(req, timeout=30) as r:
                    data = r.read()
                with open(local_path, 'wb') as f:
                    f.write(data)
                log.append(f"  ✅ {rf['name']} : téléchargé ({len(data)//1024} Ko).")
            except Exception as e:
                log.append(f"  ❌ {rf['name']} : erreur téléchargement : {e}")
    except Exception as e:
        log.append(f"[GitHub] Erreur accès API : {e}")
    return "\n".join(log)


# ─────────────────────────────────────────────────────

def open_browser():
    import time; time.sleep(1.2)
    html_files = [f for f in os.listdir('.')
                  if f.endswith('.html') and 'dashboard' in f.lower()]
    if html_files:
        url = f"http://localhost:{PORT}/{html_files[0]}"
        print(f"\n  Ouverture : {url}")
        webbrowser.open(url)


def kill_port(port):
    """Libère le port si déjà occupé"""
    try:
        result = subprocess.run(
            f'netstat -ano | findstr :{port}',
            shell=True, capture_output=True, text=True
        )
        for line in result.stdout.strip().split('\n'):
            if f':{port}' in line and 'LISTENING' in line:
                pid = line.strip().split()[-1]
                subprocess.run(f'taskkill /f /pid {pid}', shell=True,
                             capture_output=True)
                print(f"  Port {port} libéré (PID {pid})")
    except:
        pass


def main():
    kill_port(PORT)
    print("=" * 55)
    print("   DASHBOARD KPI — SERVEUR LOCAL")
    print("=" * 55)
    print(f"\n  Dossier : {os.getcwd()}")
    print(f"  Port    : {PORT}")
    print(f"\n  Le dashboard va s'ouvrir dans votre navigateur...")
    print(f"\n  Pour arreter : fermer cette fenetre")
    print("=" * 55)
    threading.Thread(target=open_browser, daemon=True).start()
    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("", PORT), DashboardHandler) as httpd:
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\n\n  Serveur arrete.")


if __name__ == "__main__":
    main()
